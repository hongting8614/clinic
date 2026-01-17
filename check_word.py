#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""验证 Word 文档内容"""

import os
from docx import Document

def check_docx(docx_file):
    """检查 Word 文档内容"""
    try:
        doc = Document(docx_file)
        
        print(f"\n{'='*60}")
        print(f"文件: {os.path.basename(docx_file)}")
        print(f"{'='*60}")
        
        # 统计信息
        para_count = len(doc.paragraphs)
        table_count = len(doc.tables)
        
        print(f"段落数: {para_count}")
        print(f"表格数: {table_count}")
        
        # 显示前5个段落
        print(f"\n前5个段落预览:")
        print("-" * 60)
        for i, para in enumerate(doc.paragraphs[:5]):
            text = para.text.strip()
            if text:
                print(f"{i+1}. {text[:80]}{'...' if len(text) > 80 else ''}")
        
        # 检查是否有乱码
        has_garbled = False
        for para in doc.paragraphs[:20]:
            text = para.text
            # 检查是否包含常见的乱码字符
            if any(char in text for char in ['�', '锘�', '銆�']):
                has_garbled = True
                break
        
        if has_garbled:
            print("\n⚠️  警告: 检测到可能的乱码")
        else:
            print("\n✅ 内容正常")
        
        return True
        
    except Exception as e:
        print(f"\n❌ 错误: {e}")
        return False

def main():
    """主函数"""
    print("="*60)
    print("Word 文档内容验证")
    print("="*60)
    
    docs = [
        "用户文档导航中心.docx",
        "5分钟快速入门.docx",
        "项目联系人信息表.docx",
    ]
    
    output_dir = "Word文档"
    
    for doc_name in docs:
        doc_path = os.path.join(output_dir, doc_name)
        if os.path.exists(doc_path):
            check_docx(doc_path)
        else:
            print(f"\n文件不存在: {doc_name}")
    
    print("\n" + "="*60)
    print("验证完成")
    print("="*60)

if __name__ == "__main__":
    main()

