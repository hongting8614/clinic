#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""简单的 Markdown 转 Word 工具"""

import os
import sys

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("❌ 缺少 python-docx 库")
    print("请运行: pip install python-docx")
    sys.exit(1)

def simple_convert(md_file, docx_file):
    """简单转换 Markdown 到 Word"""
    doc = Document()
    
    # 读取文件
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 简单处理：按行添加
    for line in content.split('\n'):
        line = line.strip()
        
        if not line:
            continue
            
        # 标题
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('#').strip()
            doc.add_heading(text, level=min(level, 3))
        # 普通文本
        else:
            # 移除 markdown 标记
            text = line.replace('**', '').replace('*', '').replace('`', '')
            if text and not text.startswith('---'):
                doc.add_paragraph(text)
    
    doc.save(docx_file)
    print(f"OK: {os.path.basename(docx_file)}")

# 文档列表
docs = [
    ("📚 用户文档导航中心.md", "用户文档导航中心.docx"),
    ("⚡ 5分钟快速入门.md", "5分钟快速入门.docx"),
    ("📖 操作流程图解.md", "操作流程图解.docx"),
    ("📱 小程序功能及用法手册.md", "小程序功能及用法手册.docx"),
    ("❓ 常见问题速查手册.md", "常见问题速查手册.docx"),
    ("📋 角色权限完整对照表.md", "角色权限完整对照表.docx"),
    ("📱 微信小程序体验版使用指南.md", "微信小程序体验版使用指南.docx"),
    ("✅ 项目部署与验证检查清单.md", "项目部署与验证检查清单.docx"),
    ("👥 项目联系人信息表.md", "项目联系人信息表.docx"),
    ("📚 项目文档总览.md", "项目文档总览.docx"),
]

# 创建输出目录
os.makedirs("Word文档", exist_ok=True)

print("开始转换...\n")
for md, docx in docs:
    if os.path.exists(md):
        try:
            simple_convert(md, os.path.join("Word文档", docx))
        except Exception as e:
            print(f"ERROR {md}: {e}")
    else:
        print(f"NOT FOUND: {md}")

print("\n完成！")

